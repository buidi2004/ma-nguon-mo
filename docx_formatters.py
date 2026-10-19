# -*- coding: utf-8 -*-
import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, top="D3D3D3", bottom="D3D3D3", left="D3D3D3", right="D3D3D3"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge, color in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        el = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="single" w:sz="4" w:space="0" w:color="{color}"/>')
        tcBorders.append(el)
    tcPr.append(tcBorders)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(0.8)
        
        # Header
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Báo Cáo Đồ Án Mã Nguồn Mở: Nền Tảng Xem Phim OpenFlix — ĐH Đồng Tháp")
        hrun.font.name = "Times New Roman"
        hrun.font.size = Pt(8.5)
        hrun.font.italic = True
        hrun.font.color.rgb = RGBColor(128, 128, 128)
        
        # Footer
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Trang ")
        frun.font.name = "Times New Roman"
        frun.font.size = Pt(9)
        frun.font.color.rgb = RGBColor(128, 128, 128)
        # Add page number field
        fldSimple = OxmlElement('w:fldSimple')
        fldSimple.set(qn('w:instr'), 'PAGE')
        fp._p.append(fldSimple)

def add_custom_heading(doc, text, level=1):
    style_name = f'Heading {level}' if level in [1, 2, 3] else 'Normal'
    try:
        p = doc.add_paragraph(style=style_name)
    except:
        p = doc.add_paragraph()
        
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.bold = True
    
    if level == 1:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(184, 29, 36) # Burgundy Red Netflix
    elif level == 2:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0, 51, 153) # Xanh DThU
    elif level == 3:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(30, 41, 59) # Slate Dark
    elif level == 4:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(71, 85, 105)
        run.italic = True
    return p

def add_table_of_contents_field(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.3
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run()
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'TOC \\o "1-3" \\h \\z \\u')
    r._r.append(fldSimple)


def add_body_paragraph(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.3
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.bold = True
        r_pre.font.color.rgb = RGBColor(30, 30, 30)
        
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.italic = italic
    r.font.color.rgb = RGBColor(40, 40, 40)
    return p

def add_bullet_item(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.bold = True
        r_pre.font.color.rgb = RGBColor(30, 30, 30)
        
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(40, 40, 40)
    return p

print("Core formatting functions ready.")
